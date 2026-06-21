from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from docx import Document
from openpyxl import load_workbook

from app.core.config import AppSettings, get_settings
from app.core.errors import AppError
from app.db.repositories.template_repository import TemplateRepository
from app.engine.ai_prompt_loader import analyze_ai_prompt_bindings
from app.engine.template_file_type import TemplateFileType, detect_template_file_type
from app.engine.template_parser import analyze_template_variables
from app.engine.template_requirement_service import TemplateRequirementService
from app.engine.xlsx_template_parser import analyze_xlsx_template
from app.models.settings_models import TemplateFileRecord
from app.models.template_models import TemplateRequirements
from app.storage.file_manager import save_template_file
from app.storage.paths import ensure_path_under_base
from app.storage.safe_filename import safe_template_filename


class TemplateFileService:
    def __init__(self, repository: TemplateRepository, settings: AppSettings | None = None) -> None:
        self.repository = repository
        self.settings = settings or get_settings()

    def list_template_files(self) -> list[TemplateFileRecord]:
        return self.repository.list_template_files()

    def create_template_file(
        self,
        template_name: str,
        description: str,
        original_filename: str,
        file_bytes: bytes,
    ) -> TemplateFileRecord:
        cleaned_name = template_name.strip()
        if not cleaned_name:
            raise AppError("BAD_REQUEST", "模板名称不能为空", 400)
        file_type = self._get_file_type(original_filename)
        self._validate_template_file(file_type, file_bytes)

        now = self._now()
        placeholder_file = f"pending_{uuid4().hex}{file_type.suffix}"
        template_id = self.repository.create_template_file_placeholder(
            template_name=cleaned_name,
            original_filename=original_filename,
            description=description.strip(),
            now=now,
            placeholder_file=placeholder_file,
        )
        stored_filename = safe_template_filename(template_id, file_type.suffix)
        template_path = f"templates/{stored_filename}"
        try:
            save_template_file(template_id, file_bytes, self.settings, file_type.suffix)
            self.repository.update_template_file_path(template_id, stored_filename, template_path, self._now())
        except FileExistsError as exc:
            self.repository.delete_template_metadata(template_id)
            raise AppError("UPLOAD_FILE_INVALID", "模板系统文件名已存在，未覆盖已有文件", 400) from exc
        except Exception:
            self.repository.delete_template_metadata(template_id)
            raise

        record = self.repository.get_template_file(template_id)
        if record is None:
            raise AppError("TEMPLATE_NOT_FOUND", f"Template not found: {template_id}", 404, {"template_id": template_id})
        return record

    def get_download_path(self, template_id: int) -> tuple[Path, TemplateFileRecord]:
        record = self.repository.get_template_file(template_id)
        if record is None:
            raise AppError("TEMPLATE_NOT_FOUND", f"Template not found: {template_id}", 404, {"template_id": template_id})
        path = self._resolve_template_path(record.template_path)
        if not path.exists() or not path.is_file():
            raise AppError("FILE_NOT_FOUND", f"Template file not found: {template_id}", 404, {"template_id": template_id})
        return path, record

    def deactivate_template_file(self, template_id: int) -> TemplateFileRecord:
        if not self.repository.deactivate_template(template_id, self._now()):
            raise AppError("TEMPLATE_NOT_FOUND", f"Template not found: {template_id}", 404, {"template_id": template_id})
        return self._required_record(template_id)

    def activate_template_file(self, template_id: int) -> TemplateFileRecord:
        if not self.repository.activate_template(template_id, self._now()):
            raise AppError("TEMPLATE_NOT_FOUND", f"Template not found: {template_id}", 404, {"template_id": template_id})
        return self._required_record(template_id)

    def replace_template_file(self, template_id: int, original_filename: str, file_bytes: bytes) -> TemplateFileRecord:
        record = self._required_record(template_id)
        new_type = self._get_file_type(original_filename)
        current_type = self._get_file_type(record.template_file)
        if new_type is not current_type:
            raise AppError("UPLOAD_FILE_INVALID", "替换模板必须与当前模板文件类型一致", 400)

        self._validate_template_file(new_type, file_bytes)
        temp_path = self._write_temp_template_file(template_id, file_bytes, new_type)
        try:
            try:
                requirements = TemplateRequirementService(self.repository.connection).get_template_requirements(template_id)
            except ValueError as exc:
                raise AppError(
                    "VALIDATION_FAILED",
                    "Template replacement pre-validation failed",
                    400,
                    {"items": [{"code": "template_requirements_invalid", "message": str(exc)}]},
                ) from exc
            self._validate_template_bindings(temp_path, requirements)
            self._resolve_template_path(record.template_path).write_bytes(file_bytes)
            if not self.repository.update_template_original_filename(template_id, original_filename, self._now()):
                raise AppError("TEMPLATE_NOT_FOUND", f"Template not found: {template_id}", 404)
        finally:
            temp_path.unlink(missing_ok=True)
        return self._required_record(template_id)

    def _validate_template_file(self, file_type: TemplateFileType, file_bytes: bytes) -> None:
        try:
            if file_type is TemplateFileType.XLSX:
                workbook = load_workbook(BytesIO(file_bytes), data_only=False)
                if not workbook.worksheets:
                    raise ValueError("workbook has no worksheets")
                workbook.close()
            else:
                Document(BytesIO(file_bytes))
        except Exception as exc:
            raise AppError("UPLOAD_FILE_INVALID", f"上传文件不是有效的 {file_type.suffix} 模板文件", 400) from exc

    def _validate_template_bindings(self, template_path: Path, requirements: TemplateRequirements) -> None:
        errors: list[dict[str, object]] = []
        allowed_tables = {table.table_name for table in requirements.required_tables}
        fields_by_table: dict[str, set[str]] = {}
        for field in requirements.fields:
            fields_by_table.setdefault(field.table_name, set()).add(field.field_name)

        file_type = detect_template_file_type(template_path)
        if file_type is TemplateFileType.XLSX:
            analysis = analyze_xlsx_template(template_path, requirements.fields)
            references = list(analysis.references)
            missing_variables = list(analysis.missing_variables)
            for code, message, coordinate in analysis.issues:
                errors.append({"code": code, "message": message, "detail": {"coordinate": coordinate}})
        else:
            references, _mapping, missing_variables = analyze_template_variables(template_path, requirements.fields)

        for missing in missing_variables:
            code = "template_variable_legacy_data_alias" if missing.table_name == "data" else "missing_field"
            errors.append({
                "code": code,
                "message": f"Template references unknown variable: {missing.original_var_path}",
                "table_name": missing.table_name,
                "field_name": missing.field_name,
                "detail": {"reason": missing.reason},
            })
        for reference in references:
            original_path = reference.original_variable_path or reference.variable_path
            if reference.table_name == "data":
                errors.append({"code": "template_variable_legacy_data_alias", "message": f"Template variable must use table.field instead of data.field: {original_path}"})
            elif reference.table_name not in allowed_tables:
                errors.append({"code": "template_variable_table_not_configured", "message": f"Template references table not configured for this template: {reference.table_name}"})
            elif reference.field_name not in fields_by_table.get(reference.table_name, set()):
                errors.append({"code": "missing_field", "message": f"Template references field not found in schema: {reference.variable_path}"})

        if file_type is TemplateFileType.DOCX:
            try:
                ai_analysis = analyze_ai_prompt_bindings(template_path)
            except Exception as exc:
                errors.append({"code": "ai_block_parse_failed", "message": f"Failed to parse AI block bindings: {exc}"})
            else:
                for issue in ai_analysis.issues:
                    errors.append({"code": issue.code, "message": issue.message})

        if errors:
            raise AppError("VALIDATION_FAILED", "Template replacement pre-validation failed", 400, {"items": errors})

    def _write_temp_template_file(self, template_id: int, file_bytes: bytes, file_type: TemplateFileType) -> Path:
        self.settings.temp_dir.mkdir(parents=True, exist_ok=True)
        temp_path = ensure_path_under_base(
            self.settings.temp_dir / f"template_replace_{template_id}_{uuid4().hex}{file_type.suffix}",
            self.settings.temp_dir,
        )
        temp_path.write_bytes(file_bytes)
        return temp_path

    def _required_record(self, template_id: int) -> TemplateFileRecord:
        record = self.repository.get_template_file(template_id)
        if record is None:
            raise AppError("TEMPLATE_NOT_FOUND", f"Template not found: {template_id}", 404, {"template_id": template_id})
        return record

    def _get_file_type(self, filename: str) -> TemplateFileType:
        try:
            return detect_template_file_type(filename)
        except ValueError as exc:
            raise AppError("UPLOAD_FILE_INVALID", "仅支持上传 .docx 或 .xlsx 模板文件", 400) from exc

    def _resolve_template_path(self, path_value: str) -> Path:
        path = Path(path_value)
        if not path.is_absolute():
            path = self.settings.project_root / path
        return ensure_path_under_base(path, self.settings.templates_dir)

    def _now(self) -> str:
        return datetime.now(UTC).isoformat()
